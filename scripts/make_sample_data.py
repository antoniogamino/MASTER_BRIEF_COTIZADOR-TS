"""Genera data/sample/COST_MASTER_SAMPLE.xlsx y HISTORICAL_PURCHASES_SAMPLE.xlsx.

*** SAMPLE DATA *** — valores ficticios para pruebas. NO son tarifas ni precios reales de Technology & Steel.
"""

from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).resolve().parents[1] / "data" / "sample"

SAMPLE = True
MAT_COLS = ["price_id", "grade", "form", "thickness_min_mm", "thickness_max_mm", "price_per_kg", "currency", "price_date",
            "supplier", "source_level", "source_type", "reference", "adjustment_factor", "is_sample"]
MATERIALS = [
    ["S-MAT-001", "AISI 304", "sheet", 0.5, 3.0, 92.50, "MXN", date(2026, 9, 20), "Proveedor Ejemplo A", "A", "supplier_quote", "COT-EJ-1001", 1, SAMPLE],
    ["S-MAT-002", "SS304", "sheet", 0.5, 3.0, 78.00, "MXN", date(2026, 3, 1), None, "D", "parametric", None, 1, SAMPLE],
    ["S-MAT-003", "AISI 304", "plate", 3.01, 25.0, 4.60, "USD", date(2026, 8, 25), "Proveedor Ejemplo B", "A", "supplier_quote", "COT-EJ-1002", 1, SAMPLE],
    ["S-MAT-004", "316L", "round_bar", None, None, 6.10, "USD", date(2026, 7, 2), "Proveedor Ejemplo C", "B", "purchase_order", "OC-EJ-4532", 1, SAMPLE],
    ["S-MAT-005", "AISI 316L", "sheet", 0.5, 3.0, 128.00, "MXN", date(2025, 12, 15), "Proveedor Ejemplo A", "C", "historical", "OC-EJ-3100", 1.08, SAMPLE],
    ["S-MAT-006", "AISI 304", "sanitary_tube", None, None, 135.00, "MXN", date(2026, 9, 1), "Proveedor Ejemplo D", "A", "supplier_quote", "COT-EJ-1010", 1, SAMPLE],
]
PROC_COLS = ["rate_key", "description", "hourly_rate", "currency", "rate_type", "minimum_charge", "effective_date", "estimated", "is_sample"]
PROCESSES = [
    ["laser_cutting", "Láser fibra (máquina + operador)", 1450, "MXN", "machine_plus_labor", 350, date(2026, 6, 1), False, SAMPLE],
    ["bending", "Prensa dobladora + operador", 780, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["rolling", "Roladora + operador", 650, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["welding", "Soldador TIG (cargado) + equipo", 520, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["surface_polishing", "Pulidor (cargado) + equipo", 430, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["weld_finishing", "Acabado de soldadura", 430, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["cnc_lathe", "Torno CNC + operador", 950, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["cnc_mill", "Centro de maquinado + operador", 1150, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["manual_lathe", "Torno convencional", 560, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["manual_mill", "Fresadora convencional", 600, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["drill_press", "Taladro de banco", 380, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["passivation", "Pasivado (químico + mano de obra)", 480, "MXN", "labor_only", None, date(2026, 6, 1), True, SAMPLE],
    ["electropolishing", "Electropulido", 900, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["inspection", "Inspector de calidad", 450, "MXN", "labor_only", None, date(2026, 6, 1), False, SAMPLE],
    ["general_labor", "Ayudante general / ensamble", 280, "MXN", "labor_only", None, date(2026, 6, 1), False, SAMPLE],
    ["saw_cutting", "Sierra cinta", 320, "MXN", "machine_plus_labor", None, date(2026, 6, 1), False, SAMPLE],
    ["plasma_cutting", "Plasma CNC", 900, "MXN", "machine_plus_labor", None, date(2026, 6, 1), True, SAMPLE],
    ["engineering", "Ingeniería / CAD / programación", 650, "MXN", "labor_only", None, date(2026, 6, 1), False, SAMPLE],
]
LABOR_COLS = ["labor_key", "salary_cost", "benefits", "burden", "indirect", "productive_hours_per_month", "currency", "is_sample"]
LABOR = [
    ["welder_tig", 24000, 7200, 3500, 2500, 176, "MXN", SAMPLE],
    ["polisher", 18000, 5400, 2600, 2000, 176, "MXN", SAMPLE],
]
ITEM_COLS = ["item_key", "description", "unit", "unit_cost", "currency", "price_date", "supplier", "is_sample"]
CONSUMABLES = [
    ["argon", "Argón industrial", "L", 0.22, "MXN", date(2026, 8, 1), "Gases Ejemplo", SAMPLE],
    ["filler_rod", "Varilla ER308L/ER316L", "kg", 420, "MXN", date(2026, 8, 1), "Soldaduras Ejemplo", SAMPLE],
    ["tungsten", "Electrodo tungsteno 2% Th", "ea", 38, "MXN", date(2026, 8, 1), "Soldaduras Ejemplo", SAMPLE],
]
PACKAGING = [
    ["cardboard_box", "Caja de cartón doble corrugado", "ea", 85, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["wood_crate_m3", "Caja de madera (por m3 exterior)", "m3", 4200, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["pallet", "Tarima", "ea", 280, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["plastic_bag", "Bolsa de polietileno", "ea", 6, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["desiccant", "Desecante", "ea", 12, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["foam", "Espuma de protección", "ea", 25, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["nimf15", "Tratamiento NIMF-15", "ea", 350, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["surface_protection", "Película protectora", "ea", 18, "MXN", date(2026, 7, 1), None, SAMPLE],
]
QUALITY = [
    ["cert_en10204_3_1", "Certificado de material EN 10204 3.1", "lot", 450, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["dimensional_report", "Reporte dimensional", "lot", 900, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["passivation_certificate", "Certificado de pasivado", "lot", 600, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["dye_penetrant", "Líquidos penetrantes (por pieza)", "ea", 350, "MXN", date(2026, 7, 1), None, SAMPLE],
    ["roughness_measurement", "Medición de rugosidad", "lot", 700, "MXN", date(2026, 7, 1), None, SAMPLE],
]
LOGISTICS = [
    ["freight_local", "Flete local ZMG", "lot", 1200, "MXN", date(2026, 7, 1), None, SAMPLE],
]
FX_COLS = ["base", "quote", "reference_fx", "current_fx", "protected_fx", "fx_date", "fx_source", "is_sample"]
FX = [
    ["USD", "MXN", 18.20, 18.45, 19.00, date(2026, 10, 1), "EJEMPLO (no oficial)", SAMPLE],
    ["EUR", "MXN", 20.10, 20.30, 21.00, date(2026, 10, 1), "EJEMPLO (no oficial)", SAMPLE],
]
RISK_COLS = ["factor_key", "weight"]
HIST_COLS = ["purchase_id", "date", "supplier", "supplier_part_number", "description", "material_grade", "material_form",
             "dimensions", "quantity", "unit", "currency", "unit_price", "total_price", "purchase_order", "project", "notes",
             "thickness_mm", "adjustment_factor", "is_sample"]
HIST = [
    ["H-001", date(2026, 9, 10), "Proveedor Ejemplo A", "LAM-304-16", "Lámina 304 cal 16 2B", "SS 304", "sheet", "1219x2438x1.5", 120, "kg", "MXN", 95.10, 11412, "OC-EJ-5001", "EJ-01", None, 1.5, None, SAMPLE],
    ["H-002", date(2025, 11, 3), "Proveedor Ejemplo A", "LAM-304-16", "Lámina 304 cal 16 2B", "AISI304", "sheet", "1219x2438x1.5", 80, "kg", "MXN", 84.00, 6720, "OC-EJ-4100", "EJ-02", None, 1.5, 1.07, SAMPLE],
    ["H-003", date(2026, 8, 15), "Proveedor Ejemplo C", "RB-316L-12", "Barra redonda 316L 1/2\"", "316L", "round_bar", "Ø12.7x3000", 30, "lb", "USD", 2.90, 87, "OC-EJ-5010", "EJ-03", None, None, None, SAMPLE],
    ["H-004", date(2026, 9, 1), "Ferretería Ejemplo", None, "Discos flap", None, None, None, 20, "ea", "MXN", 45, 900, "OC-EJ-5020", None, None, None, None, SAMPLE],
]


def _sheet(wb: Workbook, name: str, cols: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(name)
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for r in rows:
        ws.append(r)
    for i, col in enumerate(cols, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, len(col) + 2)


def main(out: Path = OUT) -> None:
    out.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "*** SAMPLE DATA *** Valores ficticios para pruebas del sistema. NO son datos reales de Technology & Steel."
    readme["A1"].font = Font(bold=True, color="C00000")
    _sheet(wb, "Materials", MAT_COLS, MATERIALS)
    _sheet(wb, "Processes", PROC_COLS, PROCESSES)
    _sheet(wb, "Labor", LABOR_COLS, LABOR)
    _sheet(wb, "Consumables", ITEM_COLS, CONSUMABLES)
    _sheet(wb, "Packaging", ITEM_COLS, PACKAGING)
    _sheet(wb, "Quality", ITEM_COLS, QUALITY)
    _sheet(wb, "Logistics", ITEM_COLS, LOGISTICS)
    _sheet(wb, "FX", FX_COLS, FX)
    _sheet(wb, "RiskParameters", RISK_COLS, [])
    wb.save(out / "COST_MASTER_SAMPLE.xlsx")

    wb2 = Workbook()
    wb2.remove(wb2.active)
    _sheet(wb2, "Purchases", HIST_COLS, HIST)
    wb2.save(out / "HISTORICAL_PURCHASES_SAMPLE.xlsx")
    print(f"OK -> {out}")


if __name__ == "__main__":
    main(Path(sys.argv[1]) if len(sys.argv) > 1 else OUT)
