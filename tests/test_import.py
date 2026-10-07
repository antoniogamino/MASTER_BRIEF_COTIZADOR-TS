from datetime import date
from decimal import Decimal

from openpyxl import Workbook, load_workbook

from ts_quoter.importers.excel import import_cost_master, import_historical_purchases, merge_prices


def test_sample_master_imports(master):
    assert master.is_sample
    assert master.rate("welding").hourly_rate == Decimal("520")
    m = next(p for p in master.materials if p.price_id == "S-MAT-002")
    assert m.grade == "AISI 304" and m.grade_raw == "SS304"  # normalizado, raw conservado


def test_missing_column_is_specific_error(tmp_path, sample_xlsx, rules):
    wb = load_workbook(sample_xlsx / "COST_MASTER_SAMPLE.xlsx")
    ws = wb["Processes"]
    ws.delete_cols(3)  # hourly_rate
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    m, rep = import_cost_master(p, rules)
    assert m is None and not rep.ok
    errs = [e for s in rep.sheets if s.sheet == "Processes" for e in s.errors]
    assert any("hourly_rate" in e for e in errs)


def test_bad_row_skipped_with_reason(tmp_path, sample_xlsx, rules):
    wb = load_workbook(sample_xlsx / "COST_MASTER_SAMPLE.xlsx")
    ws = wb["Materials"]
    ws.cell(2, 7, "PESOS")  # currency inválida
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    m, rep = import_cost_master(p, rules)
    sr = next(s for s in rep.sheets if s.sheet == "Materials")
    assert sr.skipped == 1 and any("fila 2" in e and "moneda" in e for e in sr.errors)
    assert m is not None  # el resto se importa; el reporte lo muestra


def test_sample_and_production_cannot_mix(tmp_path, sample_xlsx, rules):
    wb = load_workbook(sample_xlsx / "COST_MASTER_SAMPLE.xlsx")
    wb["Processes"].cell(2, 9, False)  # una fila "producción"
    p = tmp_path / "mixed.xlsx"
    wb.save(p)
    m, rep = import_cost_master(p, rules)
    assert m is None and any("SAMPLE" in f for f in rep.fatal)


def test_missing_required_sheet(tmp_path, rules):
    wb = Workbook()
    wb.active.title = "Materials"
    p = tmp_path / "x.xlsx"
    wb.save(p)
    m, rep = import_cost_master(p, rules)
    assert m is None and any("Processes" in f for f in rep.fatal)


def test_historical_purchases(sample_xlsx, rules, master):
    prices, rep = import_historical_purchases(sample_xlsx / "HISTORICAL_PURCHASES_SAMPLE.xlsx", rules, date(2026, 10, 6))
    by = {p.price_id: p for p in prices}
    assert by["HP-H-001"].source_level.value == "B" and by["HP-H-001"].grade == "AISI 304"
    assert by["HP-H-002"].source_level.value == "C" and by["HP-H-002"].adjustment_factor == Decimal("1.07")
    assert by["HP-H-003"].price_per_kg == Decimal("2.90") / Decimal("0.45359237")  # lb -> kg
    assert "HP-H-004" not in by  # no es material
    merged = merge_prices(master, prices, "hist")
    assert merged.version != master.version and len(master.materials) == 6  # original intacto
