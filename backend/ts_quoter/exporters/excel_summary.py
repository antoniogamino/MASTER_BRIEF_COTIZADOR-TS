"""Internal Costing Report (Excel). NO es cotización para cliente (sec. 103)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..pipeline import PipelineResult

HDR = PatternFill("solid", fgColor="1F4E78")
WARN = PatternFill("solid", fgColor="FCE4D6")


def _hdr(ws, row: list[str]) -> None:
    ws.append(row)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HDR


def export_excel(res: PipelineResult, path: str | Path) -> Path:
    p, c, r, q = res.job.project, res.cost, res.risk, res.quote
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "INTERNAL COSTING REPORT — USO INTERNO, NO ENVIAR AL CLIENTE"
    ws["A1"].font = Font(bold=True, color="C00000", size=13)
    if c.uses_sample_data:
        ws["A2"] = "*** SAMPLE DATA *** — no válido para cotizar"
        ws["A2"].font = Font(bold=True, color="C00000")
    rows = [
        ("Project", p.project_id), ("Customer", p.customer.value), ("Drawing", p.drawing_number.value),
        ("Revision", p.drawing_revision.value), ("Quantity", p.quantity.value), ("Currency", c.currency),
        ("Status", q.status), ("CEI", f"{r.cei} — {r.cei_classification}"),
        ("Overall risk", f"{r.overall_score} — {r.overall_risk}"), ("Cost coverage", f"{c.coverage.overall}%"),
        ("Rule set", c.rule_set_ref), ("Cost master", c.cost_master_version), ("As of", str(c.as_of_date)),
    ]
    for i, (k, v) in enumerate(rows, start=4):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, str(v))
    start = len(rows) + 6
    ws.cell(start, 1, "Price build-up").font = Font(bold=True, size=12)
    for i, st in enumerate(q.steps, start=start + 1):
        ws.cell(i, 1, st.label)
        ws.cell(i, 2, float(st.amount)).number_format = "#,##0.00"
        ws.cell(i, 3, st.formula)
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 80

    wl = wb.create_sheet("Cost lines")
    _hdr(wl, ["line", "component", "category", "description", "status", "one_time", "per_unit", "total@qty",
              "hours_setup", "hours_per_unit", "formula", "sources", "missing_reason"])
    qty = c.quantity or 1
    for l in c.lines:
        wl.append([l.line_id, l.component_id, l.category, l.description, l.status.value, float(l.one_time), float(l.per_unit),
                   float(l.total(qty)), float(l.hours_one_time), float(l.hours_per_unit), l.formula,
                   "; ".join(f"{s.kind}:{s.key}{' L' + s.level if s.level else ''}{' ' + str(s.price_date) if s.price_date else ''}" for s in l.sources),
                   l.missing_reason])
        if l.status.value == "missing":
            for cell in wl[wl.max_row]:
                cell.fill = WARN

    wr = wb.create_sheet("Risk")
    _hdr(wr, ["category", "factor", "weight", "severity", "points", "evidence"])
    for cat in (r.interpretation_risk, r.manufacturing_risk, r.cost_data_risk, r.commercial_risk, r.delivery_risk):
        for d in cat.drivers:
            wr.append([d.category, d.factor, float(d.weight), float(d.severity), float(d.points), d.evidence])
    wr.append([])
    _hdr(wr, ["contingency category", "base", "pct", "amount", "basis"])
    for cl in r.contingency_lines + r.rework_lines:
        wr.append([cl.category, float(cl.base_amount), float(cl.pct), float(cl.amount), cl.basis])

    wq = wb.create_sheet("Quantity breaks")
    _hdr(wq, ["qty", "base_cost", "contingency", "total_internal", "price_total", "unit_price", "rounded_unit"])
    for b in q.quantity_breaks:
        wq.append([b.quantity, float(b.base_cost), float(b.contingency), float(b.total_internal_cost),
                   float(b.suggested_price_total), float(b.suggested_unit_price), float(b.rounded_unit_price)])

    ww = wb.create_sheet("Warnings & review")
    _hdr(ww, ["severity", "cost_impact", "category", "description", "recommended_action", "question"])
    for w in q.warnings + q.stop_reasons:
        ww.append(["WARNING", "", "", w, "", ""])
    for rv in c.review_items:
        ww.append([rv.severity.value, rv.cost_impact.value, rv.category, rv.description, rv.recommended_action, rv.question])

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
