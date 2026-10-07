"""Importadores Excel con validación explícita (nunca fallar en silencio).

COST_MASTER.xlsx       -> CostMaster
HISTORICAL_PURCHASES.xlsx -> precios de material nivel B/C (merge en CostMaster nuevo, versión nueva)

Reporte: imported / skipped / warnings / errors por hoja.
"""

from __future__ import annotations

import hashlib
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import load_workbook
from pydantic import ValidationError

from ..material_engine.normalize import normalize_grade
from ..money import recency
from ..rules.loader import RuleSet
from ..schemas.cost_master import (CostMaster, FxRate, LaborRate, MaterialPrice, PriceLevel, ProcessRate,
                                   RiskParameterOverride, SimpleCostItem)
from ..schemas.manufacturing import MaterialForm
from ..units import UnitError, convert

CURRENCIES = {"MXN", "USD", "EUR"}


@dataclass
class SheetReport:
    sheet: str
    imported: int = 0
    skipped: int = 0
    warnings: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


@dataclass
class ImportReport:
    filename: str
    checksum: str
    sheets: list[SheetReport] = field(default_factory=list)
    fatal: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.fatal and not any(s.errors and s.imported == 0 for s in self.sheets if s.sheet in REQUIRED_SHEETS)

    def as_dict(self) -> dict:
        return {"filename": self.filename, "checksum": self.checksum, "ok": self.ok, "fatal": self.fatal,
                "sheets": [s.__dict__ for s in self.sheets]}

    def summary(self) -> str:
        lines = [f"Import {self.filename} [{self.checksum}] ok={self.ok}"] + [f"  FATAL: {f}" for f in self.fatal]
        for s in self.sheets:
            lines.append(f"  {s.sheet}: imported={s.imported} skipped={s.skipped} warnings={len(s.warnings)} errors={len(s.errors)}")
            lines += [f"    ERROR {e}" for e in s.errors[:10]] + [f"    warn  {w}" for w in s.warnings[:10]]
        return "\n".join(lines)


# ----------------------------------------------------------------------------- column specs
def _dec(v: Any) -> Decimal:
    if v is None or v == "":
        raise ValueError("vacío")
    try:
        return Decimal(str(v).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"no numérico: '{v}'")


def _opt_dec(v: Any) -> Optional[Decimal]:
    return None if v in (None, "") else _dec(v)


def _date(v: Any) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str) and v.strip():
        return date.fromisoformat(v.strip()[:10])
    raise ValueError(f"fecha inválida: '{v}'")


def _opt_date(v: Any) -> Optional[date]:
    return None if v in (None, "") else _date(v)


def _ccy(v: Any) -> str:
    s = str(v or "").strip().upper()
    if s not in CURRENCIES:
        raise ValueError(f"moneda no soportada: '{v}' (permitidas {sorted(CURRENCIES)})")
    return s


def _bool(v: Any) -> bool:
    return str(v).strip().lower() in ("true", "1", "yes", "si", "sí", "x") if v is not None else False


def _str(v: Any) -> str:
    s = str(v).strip() if v is not None else ""
    if not s:
        raise ValueError("vacío")
    return s


def _opt_str(v: Any) -> Optional[str]:
    s = str(v).strip() if v is not None else ""
    return s or None


Spec = dict[str, tuple[Callable[[Any], Any], bool]]  # col -> (parser, required)

ITEM_SPEC: Spec = {"item_key": (_str, True), "description": (_opt_str, False), "unit": (_str, True),
                   "unit_cost": (_dec, True), "currency": (_ccy, True), "price_date": (_opt_date, False),
                   "supplier": (_opt_str, False), "is_sample": (_bool, False)}

SHEETS: dict[str, tuple[Spec, type, str, str]] = {
    # sheet: (spec, model, master attribute, key column)
    "Materials": ({"price_id": (_str, True), "grade": (_str, True), "form": (_str, True),
                   "thickness_min_mm": (_opt_dec, False), "thickness_max_mm": (_opt_dec, False),
                   "price_per_kg": (_dec, True), "currency": (_ccy, True), "price_date": (_date, True),
                   "supplier": (_opt_str, False), "source_level": (_str, True), "source_type": (_str, True),
                   "reference": (_opt_str, False), "adjustment_factor": (_opt_dec, False), "is_sample": (_bool, False)},
                  MaterialPrice, "materials", "price_id"),
    "Processes": ({"rate_key": (_str, True), "description": (_opt_str, False), "hourly_rate": (_dec, True),
                   "currency": (_ccy, True), "rate_type": (_opt_str, False), "minimum_charge": (_opt_dec, False),
                   "effective_date": (_opt_date, False), "estimated": (_bool, False), "is_sample": (_bool, False)},
                  ProcessRate, "process_rates", "rate_key"),
    "Labor": ({"labor_key": (_str, True), "salary_cost": (_dec, True), "benefits": (_opt_dec, False),
               "burden": (_opt_dec, False), "indirect": (_opt_dec, False), "productive_hours_per_month": (_dec, True),
               "currency": (_ccy, True), "is_sample": (_bool, False)}, LaborRate, "labor", "labor_key"),
    "Consumables": (ITEM_SPEC, SimpleCostItem, "consumables", "item_key"),
    "Packaging": (ITEM_SPEC, SimpleCostItem, "packaging", "item_key"),
    "Quality": (ITEM_SPEC, SimpleCostItem, "quality", "item_key"),
    "Logistics": (ITEM_SPEC, SimpleCostItem, "logistics", "item_key"),
    "FX": ({"base": (_ccy, True), "quote": (_ccy, True), "reference_fx": (_dec, True), "current_fx": (_dec, True),
            "protected_fx": (_dec, True), "fx_date": (_date, True), "fx_source": (_str, True), "is_sample": (_bool, False)},
           FxRate, "fx", "base"),
    "RiskParameters": ({"factor_key": (_str, True), "weight": (_dec, True)}, RiskParameterOverride, "risk_parameters", "factor_key"),
}
SHEETS["Machines"] = SHEETS["Processes"][0], ProcessRate, "process_rates", "rate_key"
REQUIRED_SHEETS = {"Materials", "Processes"}


def _checksum(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:16]


def _read_sheet(ws, spec: Spec, model: type, key_col: str, rep: SheetReport) -> list:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        rep.errors.append("hoja vacía")
        return []
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    missing = [c for c, (_, req) in spec.items() if req and c not in header]
    unknown = [h for h in header if h and h not in spec]
    if missing:
        rep.errors.append(f"columnas obligatorias faltantes: {missing}; columnas encontradas: {header}")
        return []
    if unknown:
        rep.warnings.append(f"columnas ignoradas (no reconocidas): {unknown}")
    out, seen = [], {}
    for r_i, row in enumerate(rows[1:], start=2):
        if all(v in (None, "") for v in row):
            continue
        rec = dict(zip(header, row))
        data, errs = {}, []
        for col, (parser, req) in spec.items():
            v = rec.get(col)
            if v in (None, "") and not req:
                parsed = parser(v) if parser in (_bool,) else None
                if parsed is not None:
                    data[col] = parsed
                continue
            try:
                data[col] = parser(v)
            except ValueError as e:
                errs.append(f"{col}: {e}")
        if errs:
            rep.errors.append(f"fila {r_i}: " + "; ".join(errs))
            rep.skipped += 1
            continue
        try:
            obj = model(**data)
        except ValidationError as e:
            rep.errors.append(f"fila {r_i}: {e.errors()[0]['loc']} {e.errors()[0]['msg']}")
            rep.skipped += 1
            continue
        k = (data.get(key_col), data.get("quote")) if model is FxRate else data.get(key_col)
        if k in seen:
            rep.warnings.append(f"fila {r_i}: posible duplicado de '{k}' (fila {seen[k]}); se conservan ambos, se usa el primero")
        else:
            seen[k] = r_i
        out.append(obj)
        rep.imported += 1
    return out


def _validate_materials(items: list[MaterialPrice], rules: Optional[RuleSet], rep: SheetReport) -> list[MaterialPrice]:
    forms = {f.value for f in MaterialForm}
    valid_levels = {l.value for l in PriceLevel} - {"E"}
    out = []
    for m in items:
        if m.form not in forms:
            rep.errors.append(f"{m.price_id}: forma '{m.form}' no válida ({sorted(forms)})")
            rep.imported -= 1
            rep.skipped += 1
            continue
        if m.source_level.value not in valid_levels:
            rep.errors.append(f"{m.price_id}: source_level debe ser A-D (E = missing no se importa)")
            rep.imported -= 1
            rep.skipped += 1
            continue
        if rules is not None:
            g = normalize_grade(m.grade, rules)
            if g is None:
                rep.warnings.append(f"{m.price_id}: grado '{m.grade}' sin alias; se conserva tal cual")
            elif g != m.grade:
                m = m.model_copy(update={"grade_raw": m.grade, "grade": g})
        out.append(m)
    _flag_outliers(out, rep)
    return out


def _flag_outliers(items: list[MaterialPrice], rep: SheetReport, factor: Decimal = Decimal("2")) -> None:
    """Marca possible_outlier (precio > 2× o < 0.5× la mediana de su grado/forma). Nunca elimina."""
    groups: dict[tuple, list[MaterialPrice]] = {}
    for m in items:
        groups.setdefault((m.grade, m.form, m.currency), []).append(m)
    for g, ms in groups.items():
        if len(ms) < 3:
            continue
        med = Decimal(str(statistics.median([float(m.price_per_kg) for m in ms])))
        for m in ms:
            if m.price_per_kg > med * factor or m.price_per_kg < med / factor:
                m.possible_outlier = True
                rep.warnings.append(f"{m.price_id}: possible_outlier ({m.price_per_kg} vs mediana {med} {g})")


def import_cost_master(path: str | Path, rules: Optional[RuleSet] = None) -> tuple[Optional[CostMaster], ImportReport]:
    p = Path(path)
    rep = ImportReport(filename=p.name, checksum="")
    if not p.exists():
        rep.fatal.append(f"archivo no encontrado: {p}")
        return None, rep
    rep.checksum = _checksum(p)
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception as e:  # archivo corrupto / no xlsx
        rep.fatal.append(f"no se pudo abrir como .xlsx: {e}")
        return None, rep
    data: dict[str, list] = {}
    for req in REQUIRED_SHEETS:
        if req not in wb.sheetnames:
            rep.fatal.append(f"falta hoja obligatoria '{req}' (hojas: {wb.sheetnames})")
    for name in wb.sheetnames:
        if name not in SHEETS:
            rep.sheets.append(SheetReport(sheet=name, warnings=["hoja no reconocida, ignorada"]))
            continue
        spec, model, attr, key = SHEETS[name]
        sr = SheetReport(sheet=name)
        items = _read_sheet(wb[name], spec, model, key, sr)
        if name == "Materials":
            items = _validate_materials(items, rules, sr)
        data.setdefault(attr, []).extend(items)
        rep.sheets.append(sr)
    if not rep.ok:
        return None, rep
    all_items = [i for v in data.values() for i in v if hasattr(i, "is_sample")]
    is_sample = bool(all_items) and all(i.is_sample for i in all_items)
    mixed = any(i.is_sample for i in all_items) and not is_sample
    if mixed:
        rep.fatal.append("Mezcla de SAMPLE DATA con datos de producción en el mismo archivo (prohibido, sec. 173)")
        return None, rep
    master = CostMaster(version=f"cm-{rep.checksum}", imported_at=datetime.now(timezone.utc), source_filename=p.name,
                        is_sample=is_sample, **data)
    return master, rep


# ----------------------------------------------------------------------------- historical purchases
HIST_SPEC: Spec = {"purchase_id": (_str, True), "date": (_date, True), "supplier": (_opt_str, False),
                   "supplier_part_number": (_opt_str, False), "description": (_opt_str, False),
                   "material_grade": (_opt_str, False), "material_form": (_opt_str, False), "dimensions": (_opt_str, False),
                   "quantity": (_dec, True), "unit": (_str, True), "currency": (_ccy, True), "unit_price": (_dec, True),
                   "total_price": (_opt_dec, False), "purchase_order": (_opt_str, False), "project": (_opt_str, False),
                   "notes": (_opt_str, False), "thickness_mm": (_opt_dec, False), "adjustment_factor": (_opt_dec, False),
                   "is_sample": (_bool, False)}


def import_historical_purchases(path: str | Path, rules: RuleSet, as_of: date) -> tuple[list[MaterialPrice], ImportReport]:
    """Compras reales -> precios de material. Nivel B si es reciente (≤ 'acceptable'), C si es histórica.

    Nivel C usa adjustment_factor (columna); si falta se registra 1.0 con advertencia.
    """
    p = Path(path)
    rep = ImportReport(filename=p.name, checksum=_checksum(p) if p.exists() else "")
    if not p.exists():
        rep.fatal.append(f"archivo no encontrado: {p}")
        return [], rep
    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    sr = SheetReport(sheet=wb.sheetnames[0])
    rep.sheets.append(sr)

    class _Row(dict):
        def __init__(self, **kw: Any) -> None:
            super().__init__(**kw)

    rows = _read_sheet(ws, HIST_SPEC, _Row, "purchase_id", sr)
    forms = {f.value for f in MaterialForm}
    out: list[MaterialPrice] = []
    for r in rows:
        pid = r["purchase_id"]
        if not r.get("material_grade"):
            sr.warnings.append(f"{pid}: sin material_grade (no es compra de material), omitido")
            sr.skipped += 1
            continue
        grade = normalize_grade(r["material_grade"], rules)
        if grade is None:
            sr.warnings.append(f"{pid}: grado '{r['material_grade']}' sin alias; omitido")
            sr.skipped += 1
            continue
        form = (r.get("material_form") or "").strip().lower()
        if form not in forms:
            sr.warnings.append(f"{pid}: forma '{form}' no válida; omitido")
            sr.skipped += 1
            continue
        try:
            per_kg = r["unit_price"] / convert(1, r["unit"], "kg")
        except UnitError:
            sr.warnings.append(f"{pid}: unidad '{r['unit']}' no convertible a kg; omitido (no se adivina el peso)")
            sr.skipped += 1
            continue
        age, rec = recency(r["date"], as_of, rules)
        level = PriceLevel.B if rec in ("fresh", "acceptable") else PriceLevel.C
        adj = r.get("adjustment_factor") or Decimal("1")
        if level == PriceLevel.C and r.get("adjustment_factor") is None:
            sr.warnings.append(f"{pid}: compra histórica ({age} días) sin adjustment_factor; se usa 1.0")
        t = r.get("thickness_mm")
        out.append(MaterialPrice(price_id=f"HP-{pid}", grade=grade, grade_raw=r["material_grade"], form=form,
                                 thickness_min_mm=t, thickness_max_mm=t, price_per_kg=per_kg, currency=r["currency"],
                                 price_date=r["date"], supplier=r.get("supplier"), source_level=level,
                                 source_type="purchase_order" if level == PriceLevel.B else "historical",
                                 reference=r.get("purchase_order"), adjustment_factor=adj if level == PriceLevel.C else Decimal("1"),
                                 is_sample=r.get("is_sample", False)))
    _flag_outliers(out, sr)
    return out, rep


def merge_prices(master: CostMaster, prices: list[MaterialPrice], label: str) -> CostMaster:
    """Nuevo Cost Master (nueva versión); el anterior no se modifica (histórico nunca se sobreescribe)."""
    if any(p.is_sample for p in prices) != master.is_sample and prices:
        raise ValueError("No se pueden mezclar precios SAMPLE con un Cost Master de producción (o viceversa)")
    h = hashlib.sha256((master.version + label + "".join(p.price_id for p in prices)).encode()).hexdigest()[:16]
    return master.model_copy(update={"version": f"cm-{h}", "materials": master.materials + prices}, deep=True)
